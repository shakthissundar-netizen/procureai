"""Demo data generator for ProcureAI.

Generates realistic, professional quotation PDF and Excel files for:
- Vendor A (Apex Computech Solutions)
- Vendor B (BudgetByte Infotech)
- Vendor C (NovaCore Technologies)
- RFQ Requirements Specification (JSON)
"""

import json
from pathlib import Path
from typing import Any, Dict

import pymupdf as fitz
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def generate_rfq_json(output_dir: Path) -> Path:
    """Generate the benchmark RFQ requirements JSON."""
    rfq_data = {
        "rfq_id": "RFQ-2026-LAPTOP-500",
        "title": "Procurement of 500 Enterprise Business Laptops",
        "category": "IT Hardware",
        "quantity": 500,
        "product_specification": {
            "processor": "Intel Core i7 (12th Gen or higher) or equivalent AMD Ryzen 7",
            "ram_gb": 16,
            "storage_gb": 512,
            "storage_type": "NVMe SSD",
            "os": "Windows 11 Pro 64-bit",
        },
        "target_unit_price": 50000.0,
        "max_delivery_days": 15,
        "min_warranty_years": 3,
        "preferred_payment_terms": "Net 30 Days or higher",
        "scoring_weights": {
            "price": 0.30,
            "delivery": 0.25,
            "quality": 0.20,
            "warranty": 0.10,
            "payment": 0.10,
            "compliance": 0.05,
        },
    }
    file_path = output_dir / "rfq_requirements.json"
    with open(file_path, "w", encoding="utf-8") as f:
        json.dump(rfq_data, f, indent=2)
    return file_path


def create_quotation_pdf(
    file_path: Path,
    vendor_name: str,
    vendor_address: str,
    vendor_email: str,
    quotation_number: str,
    quotation_date: str,
    valid_until: str,
    product_name: str,
    specs: str,
    quantity: int,
    unit_price: float,
    subtotal: float,
    tax_percentage: float,
    tax_amount: float,
    additional_charges: float,
    grand_total: float,
    delivery_days: int,
    warranty_years: int,
    payment_terms: str,
    notes: str,
) -> None:
    """Generate a clean, professional quotation PDF using PyMuPDF."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)  # Standard A4: 595 x 842 pt

    # Define color palette (RGB normalized 0-1)
    primary_color = (0.08, 0.24, 0.44)  # Deep Navy
    secondary_color = (0.20, 0.25, 0.30)  # Slate Gray
    accent_bg = (0.94, 0.96, 0.98)  # Light Ice Blue
    border_color = (0.80, 0.85, 0.90)

    # 1. Header Banner
    page.draw_rect(fitz.Rect(40, 40, 555, 100), color=None, fill=primary_color)
    page.insert_text(fitz.Point(55, 70), "COMMERCIAL QUOTATION", fontsize=18, fontname="helv", color=(1, 1, 1))
    page.insert_text(fitz.Point(55, 90), f"Reference: {quotation_number}", fontsize=11, fontname="helv", color=(0.85, 0.9, 1.0))

    # 2. Vendor & Client Info Box
    page.draw_rect(fitz.Rect(40, 115, 290, 205), color=border_color, fill=accent_bg, width=0.8)
    page.insert_text(fitz.Point(50, 133), "FROM (VENDOR):", fontsize=9, fontname="helv", color=primary_color)
    page.insert_text(fitz.Point(50, 150), vendor_name, fontsize=11, fontname="helv", color=(0.1, 0.1, 0.1))
    page.insert_text(fitz.Point(50, 168), vendor_address, fontsize=9, fontname="helv", color=secondary_color)
    page.insert_text(fitz.Point(50, 185), f"Email: {vendor_email}", fontsize=9, fontname="helv", color=secondary_color)
    page.insert_text(fitz.Point(50, 198), f"Quotation #: {quotation_number}", fontsize=9, fontname="helv", color=secondary_color)

    page.draw_rect(fitz.Rect(305, 115, 555, 205), color=border_color, fill=accent_bg, width=0.8)
    page.insert_text(fitz.Point(315, 133), "TO (BUYER):", fontsize=9, fontname="helv", color=primary_color)
    page.insert_text(fitz.Point(315, 150), "ProcureAI Enterprise Systems", fontsize=11, fontname="helv", color=(0.1, 0.1, 0.1))
    page.insert_text(fitz.Point(315, 168), "Procurement Department, RFQ-2026-LAPTOP-500", fontsize=9, fontname="helv", color=secondary_color)
    page.insert_text(fitz.Point(315, 185), f"Date: {quotation_date}", fontsize=9, fontname="helv", color=secondary_color)
    page.insert_text(fitz.Point(315, 198), f"Valid Until: {valid_until}", fontsize=9, fontname="helv", color=secondary_color)

    # 3. Item Specification Table Header
    table_top = 225
    page.draw_rect(fitz.Rect(40, table_top, 555, table_top + 25), color=None, fill=primary_color)
    page.insert_text(fitz.Point(48, table_top + 17), "#", fontsize=9, fontname="helv", color=(1, 1, 1))
    page.insert_text(fitz.Point(68, table_top + 17), "Item Description & Specification", fontsize=9, fontname="helv", color=(1, 1, 1))
    page.insert_text(fitz.Point(340, table_top + 17), "Qty", fontsize=9, fontname="helv", color=(1, 1, 1))
    page.insert_text(fitz.Point(385, table_top + 17), "Unit Price (INR)", fontsize=9, fontname="helv", color=(1, 1, 1))
    page.insert_text(fitz.Point(475, table_top + 17), "Total Amount (INR)", fontsize=9, fontname="helv", color=(1, 1, 1))

    # 4. Item Row
    row_top = table_top + 25
    row_height = 65
    page.draw_rect(fitz.Rect(40, row_top, 555, row_top + row_height), color=border_color, fill=(1, 1, 1), width=0.5)
    page.insert_text(fitz.Point(48, row_top + 20), "1", fontsize=9, fontname="helv", color=(0.2, 0.2, 0.2))
    page.insert_text(fitz.Point(68, row_top + 20), product_name, fontsize=10, fontname="helv", color=(0.1, 0.1, 0.1))
    page.insert_text(fitz.Point(68, row_top + 38), specs[:65], fontsize=8, fontname="helv", color=secondary_color)
    page.insert_text(fitz.Point(68, row_top + 52), specs[65:130], fontsize=8, fontname="helv", color=secondary_color)

    page.insert_text(fitz.Point(340, row_top + 20), f"{quantity}", fontsize=9, fontname="helv", color=(0.1, 0.1, 0.1))
    page.insert_text(fitz.Point(385, row_top + 20), f"Rs. {unit_price:,.2f}", fontsize=9, fontname="helv", color=(0.1, 0.1, 0.1))
    page.insert_text(fitz.Point(475, row_top + 20), f"Rs. {subtotal:,.2f}", fontsize=9, fontname="helv", color=(0.1, 0.1, 0.1))

    # 5. Commercial & Delivery Terms Box (Left Side)
    terms_top = row_top + row_height + 15
    page.draw_rect(fitz.Rect(40, terms_top, 330, terms_top + 160), color=border_color, fill=accent_bg, width=0.8)
    page.insert_text(fitz.Point(50, terms_top + 20), "COMMERCIAL TERMS & SPECIFICATIONS:", fontsize=9, fontname="helv", color=primary_color)
    page.insert_text(fitz.Point(50, terms_top + 45), f"- Delivery Timeline: {delivery_days} Days from PO issuance", fontsize=9, fontname="helv", color=(0.1, 0.1, 0.1))
    page.insert_text(fitz.Point(50, terms_top + 70), f"- Warranty Coverage: {warranty_years} Years Comprehensive Warranty", fontsize=9, fontname="helv", color=(0.1, 0.1, 0.1))
    page.insert_text(fitz.Point(50, terms_top + 95), f"- Payment Terms: {payment_terms}", fontsize=9, fontname="helv", color=(0.1, 0.1, 0.1))
    page.insert_text(fitz.Point(50, terms_top + 120), f"- Note: {notes}", fontsize=8, fontname="helv", color=secondary_color)
    page.insert_text(fitz.Point(50, terms_top + 145), "- Compliance: ISO 9001 / OEM Authorized Partner", fontsize=8, fontname="helv", color=secondary_color)

    # 6. Financial Summary Box (Right Side)
    fin_left = 345
    fin_width = 210
    fin_top = terms_top
    fin_height = 160
    page.draw_rect(fitz.Rect(fin_left, fin_top, fin_left + fin_width, fin_top + fin_height), color=border_color, fill=(1, 1, 1), width=0.8)

    y_off = fin_top + 25
    page.insert_text(fitz.Point(fin_left + 10, y_off), "Subtotal:", fontsize=9, fontname="helv", color=secondary_color)
    page.insert_text(fitz.Point(fin_left + 120, y_off), f"Rs. {subtotal:,.2f}", fontsize=9, fontname="helv", color=(0.1, 0.1, 0.1))

    y_off += 25
    page.insert_text(fitz.Point(fin_left + 10, y_off), f"GST / Tax ({tax_percentage}%):", fontsize=9, fontname="helv", color=secondary_color)
    page.insert_text(fitz.Point(fin_left + 120, y_off), f"Rs. {tax_amount:,.2f}", fontsize=9, fontname="helv", color=(0.1, 0.1, 0.1))

    y_off += 25
    page.insert_text(fitz.Point(fin_left + 10, y_off), "Additional Charges:", fontsize=9, fontname="helv", color=secondary_color)
    page.insert_text(fitz.Point(fin_left + 120, y_off), f"Rs. {additional_charges:,.2f}", fontsize=9, fontname="helv", color=(0.1, 0.1, 0.1))

    y_off += 20
    page.draw_line(fitz.Point(fin_left + 10, y_off), fitz.Point(fin_left + fin_width - 10, y_off), color=primary_color, width=1.0)

    y_off += 25
    page.draw_rect(fitz.Rect(fin_left + 5, y_off - 15, fin_left + fin_width - 5, y_off + 20), color=None, fill=accent_bg)
    page.insert_text(fitz.Point(fin_left + 10, y_off), "Grand Total:", fontsize=10, fontname="helv", color=primary_color)
    page.insert_text(fitz.Point(fin_left + 100, y_off), f"Rs. {grand_total:,.2f}", fontsize=11, fontname="helv", color=primary_color)

    # 7. Signature & Authorization
    sig_top = terms_top + 180
    page.insert_text(fitz.Point(40, sig_top + 20), "Authorized Signatory:", fontsize=9, fontname="helv", color=secondary_color)
    page.insert_text(fitz.Point(40, sig_top + 40), f"For {vendor_name}", fontsize=10, fontname="helv", color=(0.1, 0.1, 0.1))
    page.insert_text(fitz.Point(40, sig_top + 55), "Digitally verified and stamped.", fontsize=8, fontname="helv", color=secondary_color)

    page.draw_line(fitz.Point(40, 800), fitz.Point(555, 800), color=border_color, width=0.5)
    page.insert_text(fitz.Point(40, 815), "ProcureAI Confidential - Generated for Procurement Evaluation", fontsize=8, fontname="helv", color=secondary_color)

    doc.save(str(file_path))
    doc.close()


def create_quotation_excel(
    file_path: Path,
    vendor_name: str,
    vendor_address: str,
    vendor_email: str,
    quotation_number: str,
    quotation_date: str,
    product_name: str,
    specs: str,
    quantity: int,
    unit_price: float,
    subtotal: float,
    tax_percentage: float,
    tax_amount: float,
    additional_charges: float,
    grand_total: float,
    delivery_days: int,
    warranty_years: int,
    payment_terms: str,
    notes: str,
) -> None:
    """Generate an Excel quotation workbook using openpyxl."""
    wb = openpyxl.Workbook()

    # Sheet 1: Quotation Summary
    ws = wb.active
    ws.title = "Quotation Summary"

    # Styling helpers
    header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    accent_fill = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
    white_font_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1A365D")
    bold_font = Font(name="Calibri", size=10, bold=True, color="000000")
    regular_font = Font(name="Calibri", size=10, color="000000")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E0"),
        right=Side(style="thin", color="CBD5E0"),
        top=Side(style="thin", color="CBD5E0"),
        bottom=Side(style="thin", color="CBD5E0"),
    )

    # Title
    ws.merge_cells("A1:E1")
    ws["A1"] = f"VENDOR QUOTATION - {vendor_name.upper()}"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    # Header metadata
    meta_rows = [
        ("Vendor Name", vendor_name, "Quotation Number", quotation_number),
        ("Vendor Address", vendor_address, "Quotation Date", quotation_date),
        ("Vendor Email", vendor_email, "Payment Terms", payment_terms),
        ("Delivery Period", f"{delivery_days} Days", "Warranty Period", f"{warranty_years} Years"),
    ]

    for idx, (k1, v1, k2, v2) in enumerate(meta_rows, start=3):
        ws[f"A{idx}"] = k1
        ws[f"A{idx}"].font = bold_font
        ws[f"B{idx}"] = v1
        ws[f"B{idx}"].font = regular_font

        ws[f"D{idx}"] = k2
        ws[f"D{idx}"].font = bold_font
        ws[f"E{idx}"] = v2
        ws[f"E{idx}"].font = regular_font

    # Table Header
    headers = ["Item #", "Product Name & Specification", "Quantity", "Unit Price (INR)", "Total (INR)"]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=8, column=col_idx, value=header)
        cell.font = white_font_bold
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Item Line Row
    item_row = [1, f"{product_name} - {specs}", quantity, unit_price, subtotal]
    for col_idx, val in enumerate(item_row, start=1):
        cell = ws.cell(row=9, column=col_idx, value=val)
        cell.font = regular_font
        cell.border = thin_border
        if col_idx in (1, 3):
            cell.alignment = Alignment(horizontal="center")
        elif col_idx in (4, 5):
            cell.alignment = Alignment(horizontal="right")
            cell.number_format = "#,##0.00"

    # Financial totals
    fin_rows = [
        ("Subtotal", subtotal),
        (f"Tax (GST {tax_percentage}%)", tax_amount),
        ("Additional Charges", additional_charges),
        ("Grand Total", grand_total),
    ]

    for offset, (lbl, val) in enumerate(fin_rows, start=11):
        ws.merge_cells(f"C{offset}:D{offset}")
        lbl_cell = ws.cell(row=offset, column=3, value=lbl)
        lbl_cell.font = bold_font
        lbl_cell.alignment = Alignment(horizontal="right")
        lbl_cell.fill = accent_fill

        val_cell = ws.cell(row=offset, column=5, value=val)
        val_cell.font = bold_font
        val_cell.alignment = Alignment(horizontal="right")
        val_cell.fill = accent_fill
        val_cell.number_format = "#,##0.00"

    # Column widths
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 45
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 22

    # Sheet 2: Technical Specifications
    ws_tech = wb.create_sheet(title="Technical Specs")
    ws_tech["A1"] = "Feature / Component"
    ws_tech["B1"] = "Vendor Specification"
    ws_tech["A1"].font = white_font_bold
    ws_tech["A1"].fill = header_fill
    ws_tech["B1"].font = white_font_bold
    ws_tech["B1"].fill = header_fill

    specs_table = [
        ("Product Name", product_name),
        ("Technical Specs", specs),
        ("Quantity", f"{quantity} units"),
        ("Lead Time (Delivery)", f"{delivery_days} calendar days"),
        ("Warranty Duration", f"{warranty_years} years"),
        ("Payment Terms", payment_terms),
        ("Notes & Exclusions", notes),
    ]

    for r_idx, (feat, spec) in enumerate(specs_table, start=2):
        ws_tech.cell(row=r_idx, column=1, value=feat).font = bold_font
        ws_tech.cell(row=r_idx, column=2, value=spec).font = regular_font

    ws_tech.column_dimensions["A"].width = 25
    ws_tech.column_dimensions["B"].width = 65

    wb.save(str(file_path))


def generate_all_demo_data(output_dir: Path) -> Dict[str, Path]:
    """Generate all demo quotation files according to the Master Spec."""
    output_dir.mkdir(parents=True, exist_ok=True)
    generated_files: Dict[str, Path] = {}

    # 1. RFQ JSON
    generated_files["rfq"] = generate_rfq_json(output_dir)

    # 2. Vendor A (Apex Computech Solutions)
    # - Price: ₹52,000 (Target was ₹50,000)
    # - Delivery: 10 days (Meets <= 15 days)
    # - Warranty: 3 years
    # - Payment: 30 days
    # - Math: 500 * 52,000 = 26,000,000; Tax 18% = 4,680,000; Total = 30,680,000
    vendor_a_pdf = output_dir / "vendor_a_quotation.pdf"
    vendor_a_xlsx = output_dir / "vendor_a_quotation.xlsx"
    create_quotation_pdf(
        file_path=vendor_a_pdf,
        vendor_name="Apex Computech Solutions Ltd.",
        vendor_address="Plot 42, Electronics City Phase 1, Bengaluru, KA 560100",
        vendor_email="commercial@apexcomputech.in",
        quotation_number="APX-2026-8801",
        quotation_date="2026-08-15",
        valid_until="2026-09-30",
        product_name="Enterprise Business Laptop Pro 15",
        specs="Intel Core i7-13700H (14-core), 16GB DDR5 5200MHz RAM, 512GB NVMe PCIe 4.0 SSD, 15.6 FHD IPS, Win 11 Pro",
        quantity=500,
        unit_price=52000.0,
        subtotal=26000000.0,
        tax_percentage=18.0,
        tax_amount=4680000.0,
        additional_charges=0.0,
        grand_total=30680000.0,
        delivery_days=10,
        warranty_years=3,
        payment_terms="Net 30 Days",
        notes="Standard OEM packaging, free door delivery across India within 10 days.",
    )
    create_quotation_excel(
        file_path=vendor_a_xlsx,
        vendor_name="Apex Computech Solutions Ltd.",
        vendor_address="Plot 42, Electronics City Phase 1, Bengaluru, KA 560100",
        vendor_email="commercial@apexcomputech.in",
        quotation_number="APX-2026-8801",
        quotation_date="2026-08-15",
        product_name="Enterprise Business Laptop Pro 15",
        specs="Intel Core i7-13700H, 16GB DDR5 RAM, 512GB NVMe SSD, Win 11 Pro",
        quantity=500,
        unit_price=52000.0,
        subtotal=26000000.0,
        tax_percentage=18.0,
        tax_amount=4680000.0,
        additional_charges=0.0,
        grand_total=30680000.0,
        delivery_days=10,
        warranty_years=3,
        payment_terms="Net 30 Days",
        notes="Standard OEM packaging, free door delivery across India within 10 days.",
    )
    generated_files["vendor_a_pdf"] = vendor_a_pdf
    generated_files["vendor_a_xlsx"] = vendor_a_xlsx

    # 3. Vendor B (BudgetByte Infotech)
    # - Price: ₹48,500 (Cheapest!)
    # - Delivery: 25 days (VIOLATES <= 15 days!)
    # - Warranty: 5 years (High warranty)
    # - Payment: 15-day payment
    # - Additional Charges: ₹25,000
    # - Deliberate Grand Total Mismatch:
    #   Subtotal (500 * 48,500) = ₹24,250,000
    #   Tax 18% = ₹4,365,000
    #   Addl charges = ₹25,000
    #   Calculated Total = ₹28,640,000
    #   Stated Total = ₹27,500,000 (Mismatch of ₹1,140,000!)
    vendor_b_pdf = output_dir / "vendor_b_quotation.pdf"
    vendor_b_xlsx = output_dir / "vendor_b_quotation.xlsx"
    create_quotation_pdf(
        file_path=vendor_b_pdf,
        vendor_name="BudgetByte Infotech Pvt Ltd",
        vendor_address="204 Crystal Tower, Industrial Area, Noida, UP 201301",
        vendor_email="sales@budgetbyte.co.in",
        quotation_number="BB-QUO-2026-409",
        quotation_date="2026-08-16",
        valid_until="2026-09-15",
        product_name="BudgetBook Enterprise X500",
        specs="AMD Ryzen 7 7730U (i7 Eqv 8-core), 16GB DDR4 RAM, 512GB M.2 NVMe SSD, 15.6 FHD, Windows 11 Pro",
        quantity=500,
        unit_price=48500.0,
        subtotal=24250000.0,
        tax_percentage=18.0,
        tax_amount=4365000.0,
        additional_charges=25000.0,
        grand_total=27500000.0,  # DELIBERATE MISMATCH (Calculated is 28,640,000.0)
        delivery_days=25,  # DELIVERY VIOLATION (RFQ limit is 15 days)
        warranty_years=5,
        payment_terms="Net 15 Days",
        notes="Requires transit insurance surcharge of ₹25,000. Extended manufacturing lead time.",
    )
    create_quotation_excel(
        file_path=vendor_b_xlsx,
        vendor_name="BudgetByte Infotech Pvt Ltd",
        vendor_address="204 Crystal Tower, Industrial Area, Noida, UP 201301",
        vendor_email="sales@budgetbyte.co.in",
        quotation_number="BB-QUO-2026-409",
        quotation_date="2026-08-16",
        product_name="BudgetBook Enterprise X500",
        specs="AMD Ryzen 7 7730U, 16GB DDR4 RAM, 512GB SSD, Windows 11 Pro",
        quantity=500,
        unit_price=48500.0,
        subtotal=24250000.0,
        tax_percentage=18.0,
        tax_amount=4365000.0,
        additional_charges=25000.0,
        grand_total=27500000.0,  # DELIBERATE MISMATCH
        delivery_days=25,  # DELIVERY VIOLATION
        warranty_years=5,
        payment_terms="Net 15 Days",
        notes="Requires transit insurance surcharge of ₹25,000. Extended manufacturing lead time.",
    )
    generated_files["vendor_b_pdf"] = vendor_b_pdf
    generated_files["vendor_b_xlsx"] = vendor_b_xlsx

    # 4. Vendor C (NovaCore Technologies)
    # - Price: ₹49,800 (Under target ₹50,000!)
    # - Delivery: 12 days (Compliant <= 15 days)
    # - Warranty: 3 years (Compliant >= 3 years)
    # - Payment: 45 days (Best payment terms)
    # - Math: 500 * 49,800 = 24,900,000; Tax 18% = 4,482,000; Total = 29,382,000 (Exact)
    vendor_c_pdf = output_dir / "vendor_c_quotation.pdf"
    vendor_c_xlsx = output_dir / "vendor_c_quotation.xlsx"
    create_quotation_pdf(
        file_path=vendor_c_pdf,
        vendor_name="NovaCore Technologies India",
        vendor_address="9th Floor, Cyber Hub, DLF Phase 2, Gurugram, HR 122002",
        vendor_email="enterprise@novacore.in",
        quotation_number="NC-2026-9942",
        quotation_date="2026-08-17",
        valid_until="2026-10-15",
        product_name="ThinkSmart EliteBook 15 Gen 4",
        specs="Intel Core i7-1365U vPro (10-core), 16GB DDR5 4800MHz RAM, 512GB Gen4 Performance SSD, 15.6 IPS Anti-glare, Win 11 Pro",
        quantity=500,
        unit_price=49800.0,
        subtotal=24900000.0,
        tax_percentage=18.0,
        tax_amount=4482000.0,
        additional_charges=0.0,
        grand_total=29382000.0,
        delivery_days=12,
        warranty_years=3,
        payment_terms="Net 45 Days",
        notes="Includes 3-year Next-Business-Day onsite service and free asset tagging.",
    )
    create_quotation_excel(
        file_path=vendor_c_xlsx,
        vendor_name="NovaCore Technologies India",
        vendor_address="9th Floor, Cyber Hub, DLF Phase 2, Gurugram, HR 122002",
        vendor_email="enterprise@novacore.in",
        quotation_number="NC-2026-9942",
        quotation_date="2026-08-17",
        product_name="ThinkSmart EliteBook 15 Gen 4",
        specs="Intel Core i7-1365U, 16GB DDR5 RAM, 512GB Gen4 SSD, Win 11 Pro",
        quantity=500,
        unit_price=49800.0,
        subtotal=24900000.0,
        tax_percentage=18.0,
        tax_amount=4482000.0,
        additional_charges=0.0,
        grand_total=29382000.0,
        delivery_days=12,
        warranty_years=3,
        payment_terms="Net 45 Days",
        notes="Includes 3-year Next-Business-Day onsite service and free asset tagging.",
    )
    generated_files["vendor_c_pdf"] = vendor_c_pdf
    generated_files["vendor_c_xlsx"] = vendor_c_xlsx

    return generated_files


if __name__ == "__main__":
    demo_dir = Path(__file__).resolve().parent
    files = generate_all_demo_data(demo_dir)
    print(f"Generated {len(files)} demo files successfully in {demo_dir}:")
    for k, v in files.items():
        print(f" - {k}: {v.name} ({v.stat().st_size:,} bytes)")
