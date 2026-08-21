"""Pytest fixtures for ProcureAI Document Processing and QA test suite."""

import io
from pathlib import Path
from typing import Any, Dict

import pymupdf as fitz
import openpyxl
import pytest

from document_processing.types import QuotationDict


@pytest.fixture
def demo_data_dir() -> Path:
    """Return the absolute path to the demo_data directory."""
    return Path(__file__).resolve().parent.parent / "demo_data"


@pytest.fixture
def standard_rfq() -> Dict[str, Any]:
    """Standard RFQ specification for 500 business laptops."""
    return {
        "rfq_id": "RFQ-2026-LAPTOP-500",
        "title": "Procurement of 500 Enterprise Business Laptops",
        "quantity": 500,
        "target_unit_price": 50000.0,
        "max_delivery_days": 15,
        "min_warranty_years": 3,
        "preferred_payment_terms": "Net 30 Days",
        "scoring_weights": {
            "price": 0.30,
            "delivery": 0.25,
            "quality": 0.20,
            "warranty": 0.10,
            "payment": 0.10,
            "compliance": 0.05,
        },
    }


@pytest.fixture
def single_page_pdf_bytes() -> bytes:
    """Generate in-memory single page PDF bytes."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text(fitz.Point(50, 100), "VENDOR QUOTATION", fontsize=16)
    page.insert_text(fitz.Point(50, 140), "Vendor: Acme Tech Solutions")
    page.insert_text(fitz.Point(50, 170), "Quote No: ACM-1001")
    page.insert_text(fitz.Point(50, 200), "Quantity: 100 units")
    page.insert_text(fitz.Point(50, 230), "Unit Price: Rs. 50,000")
    page.insert_text(fitz.Point(50, 260), "Delivery: 10 days")
    page.insert_text(fitz.Point(50, 290), "Warranty: 3 years")
    page.insert_text(fitz.Point(50, 320), "Payment: Net 30 Days")
    pdf_bytes = doc.write()
    doc.close()
    return pdf_bytes


@pytest.fixture
def multi_page_pdf_bytes() -> bytes:
    """Generate in-memory 3-page PDF bytes."""
    doc = fitz.open()
    for page_num in range(1, 4):
        page = doc.new_page(width=595, height=842)
        page.insert_text(fitz.Point(50, 100), f"Section {page_num}: Quotation Part {page_num}", fontsize=14)
        page.insert_text(fitz.Point(50, 150), f"Content for page {page_num} with technical clauses.", fontsize=10)
        if page_num == 3:
            page.insert_text(fitz.Point(50, 200), "Commercial Grand Total: Rs. 25,000,000", fontsize=12)
    pdf_bytes = doc.write()
    doc.close()
    return pdf_bytes


@pytest.fixture
def empty_page_pdf_bytes() -> bytes:
    """Generate in-memory PDF with a single blank page containing no text."""
    doc = fitz.open()
    doc.new_page(width=595, height=842)
    pdf_bytes = doc.write()
    doc.close()
    return pdf_bytes


@pytest.fixture
def corrupted_pdf_bytes() -> bytes:
    """Corrupted bytes pretending to be a PDF."""
    return b"%PDF-1.4\n%Corrupted trash content\x00\xff\xfe random non pdf bytes that will fail parsing"


@pytest.fixture
def sample_excel_bytes() -> bytes:
    """Generate in-memory Excel workbook with tabular and metadata content."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quote"
    ws.append(["Vendor Name", "Apex Computech Solutions"])
    ws.append(["Quotation Number", "APX-8801"])
    ws.append(["Delivery Days", "10"])
    ws.append(["Warranty Years", "3"])
    ws.append([])
    ws.append(["Item", "Quantity", "Unit Price", "Total"])
    ws.append(["Business Laptop", 500, 52000, 26000000])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def multi_sheet_excel_bytes() -> bytes:
    """Generate in-memory multi-sheet Excel workbook."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Overview"
    ws1.append(["Vendor", "NovaCore Technologies"])
    ws1.append(["Quote", "NC-9942"])

    ws2 = wb.create_sheet(title="Financials")
    ws2.append(["Item", "Qty", "Rate", "Subtotal", "Tax", "Grand Total"])
    ws2.append(["Enterprise Laptop", 500, 49800, 24900000, 4482000, 29382000])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def corrupted_excel_bytes() -> bytes:
    """Corrupted bytes pretending to be an Excel workbook."""
    return b"PK\x03\x04corrupted zip stream content that fails decompression"


@pytest.fixture
def sample_raw_ai_outputs() -> Dict[str, Dict[str, Any]]:
    """Simulated raw outputs from LLM extraction module."""
    return {
        "vendor_a": {
            "vendor_name": "Apex Computech Solutions Ltd.",
            "quotation_number": "APX-2026-8801",
            "product_name": "Enterprise Business Laptop Pro 15",
            "quantity": "500 units",
            "unit_price": "₹ 52,000.00",
            "delivery_days": "10 days",
            "warranty_years": "3 years",
            "payment_terms": "Net 30 Days",
            "tax_percentage": "18%",
            "additional_charges": "0",
            "grand_total": "₹ 30,680,000.00",
        },
        "vendor_b": {
            "vendor_name": "BudgetByte Infotech Pvt Ltd",
            "quotation_number": "BB-QUO-2026-409",
            "product_name": "BudgetBook Enterprise X500",
            "quantity": 500,
            "unit_price": "48,500 INR",
            "delivery_days": "25 days",
            "warranty_years": "5 years",
            "payment_terms": "Net 15 Days",
            "tax_percentage": 18,
            "additional_charges": "₹ 25,000",
            "grand_total": "₹ 27,500,000.00",  # Deliberate mismatch
        },
        "vendor_c": {
            "vendor_name": "NovaCore Technologies India",
            "quotation_number": "NC-2026-9942",
            "product_name": "ThinkSmart EliteBook 15 Gen 4",
            "quantity": "500",
            "unit_price": "49,800.00",
            "delivery_days": "12 Days",
            "warranty_years": "36 months",
            "payment_terms": "Net 45 Days",
            "tax_percentage": "18.0%",
            "additional_charges": "None",
            "grand_total": "29382000.00",
        },
    }
